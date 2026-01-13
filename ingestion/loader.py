from __future__ import annotations

import hashlib
import logging
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, Optional, Tuple

from .models import Document
from .normalizer import NormalizationResult, TextNormalizer
from .text_utils import normalize_text

logger = logging.getLogger(__name__)

# Pre-compiled regex patterns for performance
_LIST_PATTERN = re.compile(r"^\s*(?:[-*+]|\d+\.)\s+")


class UnsupportedDocumentError(Exception):
    """Raised when the pipeline encounters an unsupported extension."""


def slugify(value: str) -> str:
    normalized = re.sub(r"[^a-zA-Z0-9]+", "-", value.lower()).strip("-")
    return normalized or "document"


def hash_file(path: Path) -> str:
    sha = hashlib.sha256()
    with path.open("rb") as file_obj:
        for chunk in iter(lambda: file_obj.read(8192), b""):
            sha.update(chunk)
    return sha.hexdigest()


class DocumentParseError(Exception):
    """Raised when a document cannot be parsed."""

    def __init__(self, message: str, path: Path, partial_content: str = ""):
        super().__init__(message)
        self.path = path
        self.partial_content = partial_content


def _parse_pdf_date(date_str: Optional[object]) -> Optional[str]:
    """Parse PDF date string to ISO format.

    PDF dates are typically in format: D:YYYYMMDDHHmmSS+HH'mm'
    """
    if not date_str:
        return None
    if isinstance(date_str, datetime):
        return date_str.isoformat()
    if not isinstance(date_str, str):
        return None

    # Remove 'D:' prefix if present
    if date_str.startswith("D:"):
        date_str = date_str[2:]

    # Try to parse common PDF date formats
    try:
        # Basic format: YYYYMMDDHHMMSS
        if len(date_str) >= 14:
            year = date_str[0:4]
            month = date_str[4:6]
            day = date_str[6:8]
            hour = date_str[8:10]
            minute = date_str[10:12]
            second = date_str[12:14]
            return f"{year}-{month}-{day}T{hour}:{minute}:{second}"
        elif len(date_str) >= 8:
            year = date_str[0:4]
            month = date_str[4:6]
            day = date_str[6:8]
            return f"{year}-{month}-{day}"
    except (ValueError, IndexError):
        pass

    return None


def load_pdf(path: Path) -> Tuple[str, Dict[str, Any]]:
    """Load PDF file with error handling for corrupted/invalid files.

    Inserts page markers between pages in format [PAGE:N] where N is 1-indexed.
    This allows downstream chunking to track page boundaries.
    """
    from PyPDF2 import PdfReader
    from PyPDF2.errors import PdfReadError

    pages = []
    failed_pages = []
    metadata = {}

    try:
        reader = PdfReader(str(path))
        metadata["page_count"] = len(reader.pages)
        metadata["is_encrypted"] = reader.is_encrypted

        # Extract document metadata
        if reader.metadata:
            pdf_meta = reader.metadata
            if pdf_meta.title:
                metadata["title"] = pdf_meta.title
            if pdf_meta.author:
                metadata["author"] = pdf_meta.author
            if pdf_meta.creator:
                metadata["creator"] = pdf_meta.creator
            if pdf_meta.producer:
                metadata["producer"] = pdf_meta.producer
            if pdf_meta.subject:
                metadata["subject"] = pdf_meta.subject

            # Parse dates
            creation_date = _parse_pdf_date(pdf_meta.creation_date)
            if creation_date:
                metadata["creation_date"] = creation_date

            mod_date = _parse_pdf_date(pdf_meta.modification_date)
            if mod_date:
                metadata["modification_date"] = mod_date

        # Fallback: use filename as title if not in metadata
        if "title" not in metadata:
            metadata["title"] = path.stem

        if reader.is_encrypted:
            logger.warning(f"PDF is encrypted: {path}")
            metadata["parse_warning"] = "encrypted_pdf"
            return "", metadata

        for i, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
                page_text = text.strip()
                if page_text:
                    # Add page marker before content (1-indexed)
                    pages.append(f"[PAGE:{i + 1}]\n{page_text}")
                else:
                    # Empty page still gets marker for accurate page tracking
                    pages.append(f"[PAGE:{i + 1}]")
            except Exception as e:
                logger.warning(f"Failed to extract page {i + 1} from {path}: {e}")
                failed_pages.append(i + 1)
                pages.append(f"[PAGE:{i + 1}]")

        if failed_pages:
            metadata["failed_pages"] = failed_pages
            metadata["parse_warning"] = f"partial_extraction_{len(failed_pages)}_pages_failed"

    except PdfReadError as e:
        logger.error(f"Corrupted PDF file: {path} - {e}")
        raise DocumentParseError(f"Corrupted PDF: {e}", path)
    except Exception as e:
        logger.error(f"Failed to read PDF: {path} - {e}")
        raise DocumentParseError(f"PDF read error: {e}", path)

    return "\n\n".join(pages), metadata


def load_docx(path: Path) -> Tuple[str, Dict[str, Any]]:
    """Load DOCX file with table extraction and error handling."""
    import docx
    from docx.opc.exceptions import PackageNotFoundError
    import zipfile

    content_parts = []
    metadata = {}

    try:
        doc = docx.Document(str(path))

        # Extract core properties (document metadata)
        core_props = doc.core_properties
        if core_props.title:
            metadata["title"] = core_props.title
        if core_props.author:
            metadata["author"] = core_props.author
        if core_props.subject:
            metadata["subject"] = core_props.subject
        if core_props.keywords:
            metadata["keywords"] = core_props.keywords
        if core_props.category:
            metadata["category"] = core_props.category
        if core_props.comments:
            metadata["comments"] = core_props.comments
        if core_props.created:
            metadata["creation_date"] = core_props.created.isoformat()
        if core_props.modified:
            metadata["modification_date"] = core_props.modified.isoformat()
        if core_props.last_modified_by:
            metadata["last_modified_by"] = core_props.last_modified_by
        if core_props.revision is not None:
            metadata["revision"] = core_props.revision

        # Fallback: use filename as title if not in metadata
        if "title" not in metadata:
            metadata["title"] = path.stem

        # Extract paragraphs
        paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
        metadata["paragraph_count"] = len(paragraphs)
        content_parts.extend(paragraphs)

        # Extract tables
        table_count = 0
        for table in doc.tables:
            table_count += 1
            table_text = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):
                    table_text.append(" | ".join(row_cells))
            if table_text:
                content_parts.append("\n[TABLE]\n" + "\n".join(table_text) + "\n[/TABLE]")

        metadata["table_count"] = table_count

        # Extract headers and footers
        header_text = []
        footer_text = []
        for section in doc.sections:
            if section.header and section.header.paragraphs:
                for para in section.header.paragraphs:
                    if para.text.strip():
                        header_text.append(para.text.strip())
            if section.footer and section.footer.paragraphs:
                for para in section.footer.paragraphs:
                    if para.text.strip():
                        footer_text.append(para.text.strip())

        if header_text:
            metadata["has_headers"] = True
        if footer_text:
            metadata["has_footers"] = True

    except (PackageNotFoundError, zipfile.BadZipFile) as e:
        logger.error(f"Invalid or corrupted DOCX file: {path} - {e}")
        raise DocumentParseError(f"Corrupted DOCX: {e}", path)
    except Exception as e:
        logger.error(f"Failed to read DOCX: {path} - {e}")
        raise DocumentParseError(f"DOCX read error: {e}", path)

    return "\n\n".join(content_parts), metadata


def _extract_doc_with_antiword(path: Path) -> Optional[str]:
    """Extract text from legacy .doc file using antiword CLI tool.

    Args:
        path: Path to the .doc file.

    Returns:
        Extracted text if successful, None if antiword is not available.

    Raises:
        DocumentParseError: If antiword fails to process the file.
    """
    antiword_path = shutil.which("antiword")
    if not antiword_path:
        return None

    try:
        result = subprocess.run(
            [antiword_path, str(path)],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if result.returncode == 0:
            return result.stdout
        else:
            logger.warning(f"antiword failed for {path}: {result.stderr}")
            raise DocumentParseError(f"antiword error: {result.stderr}", path)
    except subprocess.TimeoutExpired:
        raise DocumentParseError("antiword timed out processing file", path)
    except FileNotFoundError:
        return None


def _extract_doc_with_win32com(path: Path) -> Optional[Tuple[str, Dict[str, Any]]]:
    """Extract text from legacy .doc file using Windows COM automation.

    Requires Microsoft Word to be installed on the system.

    Args:
        path: Path to the .doc file.

    Returns:
        Tuple of (text, metadata) if successful, None if COM is not available.

    Raises:
        DocumentParseError: If Word fails to process the file.
    """
    try:
        import win32com.client
        import pythoncom
    except ImportError:
        return None

    word = None
    doc = None
    try:
        pythoncom.CoInitialize()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False

        doc = word.Documents.Open(str(path.resolve()), ReadOnly=True)

        # Extract text from the document
        text = doc.Content.Text

        # Extract metadata from built-in properties
        metadata = {}
        try:
            props = doc.BuiltInDocumentProperties
            prop_mapping = {
                "Title": "title",
                "Author": "author",
                "Subject": "subject",
                "Keywords": "keywords",
                "Comments": "comments",
                "Last Author": "last_modified_by",
            }
            for word_prop, meta_key in prop_mapping.items():
                try:
                    value = props(word_prop).Value
                    if value:
                        metadata[meta_key] = str(value)
                except Exception:
                    pass

            # Extract dates
            try:
                created = props("Creation Date").Value
                if created:
                    metadata["creation_date"] = created.isoformat() if hasattr(created, "isoformat") else str(created)
            except Exception:
                pass

            try:
                modified = props("Last Save Time").Value
                if modified:
                    metadata["modification_date"] = modified.isoformat() if hasattr(modified, "isoformat") else str(modified)
            except Exception:
                pass
        except Exception as e:
            logger.debug(f"Could not extract metadata from .doc: {e}")

        return text, metadata

    except Exception as e:
        logger.error(f"Win32COM failed for {path}: {e}")
        raise DocumentParseError(f"Win32COM error: {e}", path)
    finally:
        if doc:
            try:
                doc.Close(False)
            except Exception:
                pass
        if word:
            try:
                word.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def load_doc(path: Path) -> Tuple[str, Dict[str, Any]]:
    """Load legacy .doc (Word 97-2003) file with multiple extraction strategies.

    Attempts extraction in the following order:
    1. antiword CLI tool (cross-platform, requires installation)
    2. Windows COM automation (Windows only, requires Microsoft Word)

    Args:
        path: Path to the .doc file.

    Returns:
        Tuple of (extracted_text, metadata_dict).

    Raises:
        DocumentParseError: If the file cannot be read by any available method.
    """
    metadata: Dict[str, Any] = {}

    # Strategy 1: Try antiword (cross-platform)
    try:
        text = _extract_doc_with_antiword(path)
        if text is not None:
            logger.debug(f"Extracted .doc with antiword: {path}")
            metadata["extraction_method"] = "antiword"
            # Fallback title from filename
            metadata["title"] = path.stem
            return text, metadata
    except DocumentParseError:
        raise

    # Strategy 2: Try Windows COM automation
    try:
        result = _extract_doc_with_win32com(path)
        if result is not None:
            text, com_metadata = result
            metadata.update(com_metadata)
            metadata["extraction_method"] = "win32com"
            # Fallback title from filename if not in metadata
            if "title" not in metadata:
                metadata["title"] = path.stem
            return text, metadata
    except DocumentParseError:
        raise

    # No extraction method available
    error_msg = (
        "Cannot extract text from legacy .doc file. "
        "Please install one of the following:\n"
        "  - antiword: Cross-platform CLI tool (https://www.winfield.demon.nl/)\n"
        "    Windows: Download from http://antiword.cjb.net/ or use chocolatey: choco install antiword\n"
        "    Linux: apt-get install antiword / yum install antiword\n"
        "    macOS: brew install antiword\n"
        "  - pywin32: Windows only, requires Microsoft Word (pip install pywin32)"
    )
    logger.error(f"No extraction method available for .doc: {path}")
    raise DocumentParseError(error_msg, path)


def _parse_yaml_frontmatter(content: str) -> Tuple[Dict[str, Any], str]:
    """Parse YAML front matter from markdown content.

    Returns:
        Tuple of (frontmatter_dict, content_without_frontmatter)
    """
    if not content.startswith("---"):
        return {}, content

    # Find the closing ---
    end_marker = content.find("\n---", 3)
    if end_marker == -1:
        return {}, content

    frontmatter_str = content[4:end_marker].strip()
    remaining_content = content[end_marker + 4:].lstrip("\n")

    # Parse YAML manually (simple key: value pairs)
    frontmatter = {}
    for line in frontmatter_str.split("\n"):
        line = line.strip()
        if not line or line.startswith("#"):
            continue

        if ":" in line:
            key, _, value = line.partition(":")
            key = key.strip().lower()
            value = value.strip()

            # Remove quotes if present
            if (value.startswith('"') and value.endswith('"')) or \
               (value.startswith("'") and value.endswith("'")):
                value = value[1:-1]

            # Handle arrays (simple format: [item1, item2])
            if value.startswith("[") and value.endswith("]"):
                items = value[1:-1].split(",")
                value = [item.strip().strip("\"'") for item in items if item.strip()]

            if value:
                frontmatter[key] = value

    return frontmatter, remaining_content


def load_markdown(path: Path) -> Tuple[str, Dict[str, Any]]:
    """Load Markdown/text file with structure-aware parsing and error handling."""
    metadata = {}

    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # Try fallback encodings (latin-1 last as it accepts any byte sequence)
        for encoding in ["cp1252", "iso-8859-1", "latin-1"]:
            try:
                content = path.read_text(encoding=encoding)
                metadata["encoding_fallback"] = encoding
                logger.warning(f"Used fallback encoding {encoding} for: {path}")
                break
            except UnicodeDecodeError:
                continue
        else:
            logger.error(f"Failed to decode file with any encoding: {path}")
            raise DocumentParseError(f"Unable to decode file", path)
    except Exception as e:
        logger.error(f"Failed to read file: {path} - {e}")
        raise DocumentParseError(f"File read error: {e}", path)

    # Parse front matter if present
    frontmatter, _ = _parse_yaml_frontmatter(content)
    if frontmatter:
        metadata["has_frontmatter"] = True
        # Extract common frontmatter fields
        if "title" in frontmatter:
            metadata["title"] = frontmatter["title"]
        if "author" in frontmatter:
            metadata["author"] = frontmatter["author"]
        if "date" in frontmatter:
            metadata["creation_date"] = frontmatter["date"]
        if "tags" in frontmatter:
            metadata["tags"] = frontmatter["tags"]
        if "description" in frontmatter:
            metadata["description"] = frontmatter["description"]
        if "categories" in frontmatter:
            metadata["categories"] = frontmatter["categories"]

    # Extract markdown structure metadata
    lines = content.split("\n")

    # Count headers by level and extract first H1 for title fallback
    headers = {"h1": 0, "h2": 0, "h3": 0, "h4": 0, "h5": 0, "h6": 0}
    first_h1 = None
    for line in lines:
        stripped = line.strip()
        if stripped.startswith("###### ") or stripped == "######":
            headers["h6"] += 1
        elif stripped.startswith("##### ") or stripped == "#####":
            headers["h5"] += 1
        elif stripped.startswith("#### ") or stripped == "####":
            headers["h4"] += 1
        elif stripped.startswith("### ") or stripped == "###":
            headers["h3"] += 1
        elif stripped.startswith("## ") or stripped == "##":
            headers["h2"] += 1
        elif stripped.startswith("# "):
            headers["h1"] += 1
            if first_h1 is None:
                first_h1 = stripped[2:].strip()
        elif stripped == "#":
            headers["h1"] += 1

    if any(headers.values()):
        metadata["headers"] = {k: v for k, v in headers.items() if v > 0}

    # Fallback: use first H1 as title, then filename
    if "title" not in metadata:
        if first_h1:
            metadata["title"] = first_h1
        else:
            metadata["title"] = path.stem

    # Detect code blocks
    code_block_count = content.count("```") // 2
    if code_block_count > 0:
        metadata["code_blocks"] = code_block_count

    # Detect lists (unordered: -, *, + and ordered: 1., 2., etc.)
    list_items = sum(1 for line in lines if _LIST_PATTERN.match(line))
    if list_items > 0:
        metadata["list_items"] = list_items

    # Detect links
    links = re.findall(r"\[([^\]]+)\]\(([^)]+)\)", content)
    if links:
        metadata["link_count"] = len(links)

    metadata["line_count"] = len(lines)

    return content, metadata


def _detect_csv_delimiter(content: str) -> str:
    """Detect the most likely delimiter in CSV content.

    Tests common delimiters and returns the one that produces the most
    consistent column count across the first few rows.

    Args:
        content: Raw CSV content as string.

    Returns:
        Detected delimiter character.
    """
    import csv

    delimiters = [",", ";", "\t", "|"]
    best_delimiter = ","
    best_score = 0

    # Get first 10 lines for analysis
    lines = content.split("\n")[:10]
    sample = "\n".join(lines)

    for delimiter in delimiters:
        try:
            reader = csv.reader(sample.split("\n"), delimiter=delimiter)
            rows = list(reader)

            if len(rows) < 2:
                continue

            # Score based on column count consistency and number of columns
            col_counts = [len(row) for row in rows if row]
            if not col_counts:
                continue

            # Prefer delimiters that give consistent column counts > 1
            avg_cols = sum(col_counts) / len(col_counts)
            consistency = 1 - (max(col_counts) - min(col_counts)) / max(max(col_counts), 1)

            score = avg_cols * consistency

            if avg_cols > 1 and score > best_score:
                best_score = score
                best_delimiter = delimiter

        except Exception:
            continue

    return best_delimiter


def load_csv(path: Path) -> Tuple[str, Dict[str, Any]]:
    """Load CSV/TSV file with delimiter detection.

    Automatically detects the delimiter (comma, semicolon, tab, or pipe)
    and extracts text in a format consistent with Excel loading.

    Args:
        path: Path to the CSV/TSV file.

    Returns:
        Tuple of (extracted_text, metadata_dict).

    Raises:
        DocumentParseError: If the file cannot be read or is malformed.
    """
    import csv

    metadata: Dict[str, Any] = {}
    rows_text: list[str] = []
    total_rows = 0
    max_columns = 0

    # Read file content with encoding fallback
    content: Optional[str] = None
    encoding_used = "utf-8"

    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # Try fallback encodings
        for encoding in ["cp1252", "iso-8859-1", "latin-1"]:
            try:
                content = path.read_text(encoding=encoding)
                encoding_used = encoding
                metadata["encoding_fallback"] = encoding
                logger.warning(f"Used fallback encoding {encoding} for: {path}")
                break
            except UnicodeDecodeError:
                continue

    if content is None:
        logger.error(f"Failed to decode CSV file with any encoding: {path}")
        raise DocumentParseError("Unable to decode CSV file", path)

    # Detect delimiter
    delimiter = _detect_csv_delimiter(content)
    metadata["detected_delimiter"] = delimiter

    # Parse CSV
    try:
        reader = csv.reader(content.split("\n"), delimiter=delimiter)

        for row in reader:
            # Skip completely empty rows
            if not any(cell.strip() for cell in row):
                continue

            total_rows += 1
            max_columns = max(max_columns, len(row))

            # Format row as pipe-separated values (consistent with Excel)
            cell_values = [cell.strip() for cell in row]
            rows_text.append(" | ".join(cell_values))

    except csv.Error as e:
        logger.error(f"CSV parsing error in {path}: {e}")
        raise DocumentParseError(f"CSV parsing error: {e}", path)
    except Exception as e:
        logger.error(f"Failed to read CSV file: {path} - {e}")
        raise DocumentParseError(f"CSV read error: {e}", path)

    # Build output text with sheet marker (consistent with Excel format)
    sheet_name = path.stem
    if rows_text:
        text = f"[SHEET:{sheet_name}]\n" + "\n".join(rows_text)
    else:
        text = f"[SHEET:{sheet_name}]"

    # Set metadata
    metadata["title"] = path.stem
    metadata["total_rows"] = total_rows
    metadata["total_columns"] = max_columns
    metadata["sheet_names"] = [sheet_name]
    metadata["sheet_count"] = 1

    return text, metadata


def load_excel(path: Path) -> Tuple[str, Dict[str, Any]]:
    """Load Excel file (.xlsx, .xls) with multi-sheet support.

    Processes all sheets in the workbook, adding [SHEET:sheet_name] markers
    between sheets (similar to [PAGE:N] for PDFs). Rows are converted to
    pipe-separated text for consistent table representation.

    Args:
        path: Path to the Excel file.

    Returns:
        Tuple of (extracted_text, metadata_dict).

    Raises:
        DocumentParseError: If the file cannot be read or is corrupted.
    """
    metadata: Dict[str, Any] = {}
    sheet_contents: list[str] = []
    total_rows = 0
    max_columns = 0

    ext = path.suffix.lower()

    try:
        if ext == ".xlsx":
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException

            try:
                wb = load_workbook(str(path), read_only=True, data_only=True)
            except InvalidFileException as e:
                logger.error(f"Invalid or corrupted Excel file: {path} - {e}")
                raise DocumentParseError(f"Corrupted Excel file: {e}", path)
            except Exception as e:
                if "password" in str(e).lower() or "encrypted" in str(e).lower():
                    logger.warning(f"Password-protected Excel file: {path}")
                    metadata["parse_warning"] = "password_protected"
                    metadata["title"] = path.stem
                    return "", metadata
                raise

            sheet_names = wb.sheetnames
            metadata["sheet_names"] = sheet_names
            metadata["sheet_count"] = len(sheet_names)

            # Extract document properties if available
            if wb.properties:
                props = wb.properties
                if props.title:
                    metadata["title"] = props.title
                if props.creator:
                    metadata["author"] = props.creator
                if props.created:
                    metadata["creation_date"] = props.created.isoformat()
                if props.modified:
                    metadata["modification_date"] = props.modified.isoformat()

            for sheet_name in sheet_names:
                ws = wb[sheet_name]
                rows_text: list[str] = []
                sheet_row_count = 0

                for row in ws.iter_rows(values_only=True):
                    # Convert cell values to strings, handling None
                    cell_values = [str(cell) if cell is not None else "" for cell in row]

                    # Skip completely empty rows
                    if not any(cell_values):
                        continue

                    sheet_row_count += 1
                    max_columns = max(max_columns, len(cell_values))
                    rows_text.append(" | ".join(cell_values))

                total_rows += sheet_row_count

                if rows_text:
                    sheet_content = f"[SHEET:{sheet_name}]\n" + "\n".join(rows_text)
                else:
                    sheet_content = f"[SHEET:{sheet_name}]"

                sheet_contents.append(sheet_content)

            wb.close()

        elif ext == ".xls":
            import xlrd

            try:
                wb = xlrd.open_workbook(str(path))
            except xlrd.biffh.XLRDError as e:
                logger.error(f"Invalid or corrupted XLS file: {path} - {e}")
                raise DocumentParseError(f"Corrupted XLS file: {e}", path)

            sheet_names = wb.sheet_names()
            metadata["sheet_names"] = sheet_names
            metadata["sheet_count"] = len(sheet_names)

            for sheet_name in sheet_names:
                ws = wb.sheet_by_name(sheet_name)
                rows_text: list[str] = []

                for row_idx in range(ws.nrows):
                    cell_values = [
                        str(ws.cell_value(row_idx, col_idx)) if ws.cell_value(row_idx, col_idx) != "" else ""
                        for col_idx in range(ws.ncols)
                    ]

                    # Skip completely empty rows
                    if not any(cell_values):
                        continue

                    total_rows += 1
                    max_columns = max(max_columns, len(cell_values))
                    rows_text.append(" | ".join(cell_values))

                if rows_text:
                    sheet_content = f"[SHEET:{sheet_name}]\n" + "\n".join(rows_text)
                else:
                    sheet_content = f"[SHEET:{sheet_name}]"

                sheet_contents.append(sheet_content)

        else:
            raise DocumentParseError(f"Unsupported Excel format: {ext}", path)

    except DocumentParseError:
        raise
    except Exception as e:
        logger.error(f"Failed to read Excel file: {path} - {e}")
        raise DocumentParseError(f"Excel read error: {e}", path)

    # Fallback: use filename as title if not in metadata
    if "title" not in metadata:
        metadata["title"] = path.stem

    metadata["total_rows"] = total_rows
    metadata["total_columns"] = max_columns

    return "\n\n".join(sheet_contents), metadata

HANDLERS: Dict[str, Callable[[Path], Tuple[str, Dict[str, Any]]]] = {
    ".pdf": load_pdf,
    ".doc": load_doc,
    ".docx": load_docx,
    ".xlsx": load_excel,
    ".xls": load_excel,
    ".csv": load_csv,
    ".tsv": load_csv,
    ".md": load_markdown,
    ".txt": load_markdown,
}


@dataclass
class DocumentLoader:
    """Document loader with optional text normalization.

    Attributes:
        input_root: Root directory for input documents.
        normalizer: Optional TextNormalizer for additional text cleaning.
    """

    input_root: Path
    normalizer: Optional[TextNormalizer] = field(default=None)

    def load(self, path: Path) -> Tuple[Document, str]:
        """Load and process a document from disk.

        Args:
            path: Path to the document file.

        Returns:
            Tuple of (Document, content_hash).

        Raises:
            UnsupportedDocumentError: If no handler exists for the file type.
        """
        ext = path.suffix.lower()
        if ext not in HANDLERS:
            raise UnsupportedDocumentError(f"No loader configured for {ext} files.")
        handler = HANDLERS[ext]
        raw_text, extra_metadata = handler(path)

        # Apply basic text normalization
        normalized_text = normalize_text(raw_text)

        # Apply advanced normalization if configured
        normalization_result: Optional[NormalizationResult] = None
        if self.normalizer is not None:
            normalization_result = self.normalizer.normalize(normalized_text)
            normalized_text = normalization_result.text

        try:
            rel_path = str(path.relative_to(self.input_root)).replace("\\", "/")
        except ValueError:
            rel_path = path.name
        stat_info = path.stat()
        metadata = {
            "source_path": str(path.resolve()),
            "relative_path": rel_path,
            "file_extension": ext.lstrip("."),
            "last_modified": datetime.fromtimestamp(stat_info.st_mtime).isoformat(),
            "size_bytes": stat_info.st_size,
        }
        metadata.update(extra_metadata or {})
        content_hash = hash_file(path)
        metadata["content_hash"] = content_hash
        metadata["ingestion_timestamp"] = datetime.utcnow().isoformat() + "Z"

        # Add normalization metadata if normalization was applied
        if normalization_result is not None:
            metadata["normalization"] = {
                "original_length": normalization_result.original_length,
                "normalized_length": normalization_result.normalized_length,
                "rules_applied": normalization_result.rules_applied,
                "removed_patterns": normalization_result.removed_patterns,
            }

        doc_id = slugify(rel_path)
        document = Document(
            doc_id=doc_id,
            path=path,
            source_type=metadata["file_extension"],
            text=normalized_text,
            metadata=metadata,
        )
        return document, content_hash


def discover_documents(input_root: Path):
    documents = []
    for ext in HANDLERS:
        documents.extend(input_root.rglob(f"*{ext}"))
    documents.sort()
    return tuple(documents)
